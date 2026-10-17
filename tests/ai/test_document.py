"""Policy document text extraction: every supported type, and every way a file is refused."""

from __future__ import annotations

import io

import docx
import pytest
from pypdf import PdfWriter

from haqwa.ai.document import MAX_BYTES, MAX_CHARS, DocumentError, extract_text

POLICY = "Refund policy\n\nA refund must be completed within 24 hours of the request."


def make_pdf(text: str) -> bytes:
    """A minimal one-page PDF with a real text layer (no PDF library needed)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    )
    return out.getvalue()


def make_docx() -> bytes:
    document = docx.Document()
    document.add_paragraph("Refund policy")
    document.add_paragraph("A refund must be completed within 24 hours of the request.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Shipping"
    table.rows[0].cells[1].text = "A cancelled order must never be shipped."
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


# ---- supported types ----------------------------------------------------------------------


@pytest.mark.parametrize("name", ["policy.txt", "POLICY.MD", "notes.md"])
def test_text_and_markdown(name) -> None:
    result = extract_text(name, POLICY.encode())
    assert result.kind == "text" and result.text == POLICY and result.filename == name


def test_utf8_bom_and_messy_whitespace_are_tidied() -> None:
    raw = "\ufeffRule one:\t  charge   once.\n\n\n\n  Rule two.\u00a0 Ship\u00a0later.  "
    result = extract_text("p.txt", raw.encode("utf-8"))
    assert result.text == "Rule one: charge once.\n\nRule two. Ship later."


def test_docx_paragraphs_and_tables() -> None:
    result = extract_text("policy.docx", make_docx())
    assert result.kind == "docx"
    assert "A refund must be completed within 24 hours of the request." in result.text
    assert "Shipping | A cancelled order must never be shipped." in result.text


def test_pdf_text_layer() -> None:
    result = extract_text("policy.pdf", make_pdf("Refunds within 24 hours."))
    assert result.kind == "pdf" and "Refunds within 24 hours." in result.text


def test_filename_path_is_reduced_to_its_name() -> None:
    assert extract_text("C:/Users/x/policy.txt", b"Charge once.").filename == "policy.txt"


# ---- refused files ------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("name", "message"),
    [("policy.doc", "Save it as .docx"), ("policy.rtf", "Unsupported"), ("policy", "(none)")],
)
def test_unsupported_types(name, message) -> None:
    with pytest.raises(DocumentError, match=message.replace("(", r"\(").replace(")", r"\)")):
        extract_text(name, b"text")


def test_empty_file() -> None:
    with pytest.raises(DocumentError, match="empty"):
        extract_text("p.txt", b"")


def test_whitespace_only_file_has_no_text() -> None:
    with pytest.raises(DocumentError, match="No text"):
        extract_text("p.txt", b"  \n\t \n")


def test_too_many_bytes() -> None:
    with pytest.raises(DocumentError, match="larger than"):
        extract_text("p.txt", b"a" * (MAX_BYTES + 1))


def test_too_many_characters() -> None:
    with pytest.raises(DocumentError, match="limit is"):
        extract_text("p.txt", b"a " * (MAX_CHARS // 2 + 10))


def test_not_utf8() -> None:
    with pytest.raises(DocumentError, match="UTF-8"):
        extract_text("p.txt", "café".encode("latin-1"))


@pytest.mark.parametrize("data", [b"not a zip", b"PK\x03\x04broken"])
def test_broken_docx(data) -> None:
    with pytest.raises(DocumentError, match=r"\.docx file could not be read"):
        extract_text("p.docx", data)


def test_broken_pdf() -> None:
    with pytest.raises(DocumentError, match="PDF could not be read"):
        extract_text("p.pdf", b"%PDF-1.4 garbage")


def test_pdf_without_text_layer_mentions_scans() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    out = io.BytesIO()
    writer.write(out)
    with pytest.raises(DocumentError, match="scanned PDF"):
        extract_text("scan.pdf", out.getvalue())


def test_password_protected_pdf() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt("secret")
    out = io.BytesIO()
    writer.write(out)
    with pytest.raises(DocumentError, match="password"):
        extract_text("locked.pdf", out.getvalue())
