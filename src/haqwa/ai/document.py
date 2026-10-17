"""Plain text from an uploaded policy document (.txt, .md, .docx, .pdf). No Gemini here.

The file is read in memory and never stored. Limits keep a large or hostile file from
using much memory or Gemini quota: at most `MAX_BYTES` in and `MAX_CHARS` of text out.
"""

from __future__ import annotations

import io
import re
from pathlib import PurePath
from typing import Literal

from pydantic import BaseModel

MAX_BYTES = 1_000_000
MAX_CHARS = 20_000
Kind = Literal["text", "docx", "pdf"]
_KINDS: dict[str, Kind] = {".txt": "text", ".md": "text", ".docx": "docx", ".pdf": "pdf"}
SUPPORTED = ", ".join(_KINDS)


class DocumentError(ValueError):
    """The file cannot be used; the message is safe to show to the policy owner."""


class DocumentText(BaseModel):
    """The text of one uploaded policy document."""

    filename: str
    kind: Kind
    text: str


def extract_text(filename: str, data: bytes) -> DocumentText:
    """Return the document's text with tidy whitespace, or raise `DocumentError`."""
    suffix = PurePath(filename).suffix.lower()
    kind = _KINDS.get(suffix)
    if kind is None:
        hint = " Save it as .docx first." if suffix == ".doc" else ""
        raise DocumentError(f"Unsupported file type {suffix or '(none)'!r}; use {SUPPORTED}.{hint}")
    if not data:
        raise DocumentError("The file is empty.")
    if len(data) > MAX_BYTES:
        raise DocumentError(f"The file is larger than {MAX_BYTES // 1_000_000} MB.")
    if kind == "text":
        raw = _text(data)
    elif kind == "docx":
        raw = _docx(data)
    else:
        raw = _pdf(data)
    text = _tidy(raw)
    if not text:
        hint = " A scanned PDF has no text layer." if kind == "pdf" else ""
        raise DocumentError(f"No text was found in the file.{hint}")
    if len(text) > MAX_CHARS:
        raise DocumentError(
            f"The document has {len(text):,} characters; the limit is {MAX_CHARS:,}. "
            "Upload only the policy section."
        )
    return DocumentText(filename=PurePath(filename).name, kind=kind, text=text)


def _text(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentError("The text file is not UTF-8. Save it as UTF-8 and try again.") from exc


def _docx(data: bytes) -> str:
    import docx  # python-docx; imported here so the rest of the package does not need it

    try:
        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
    except Exception as exc:  # any parser failure on an untrusted file is a bad file
        raise DocumentError("The .docx file could not be read.") from exc
    return "\n".join(parts)


def _pdf(data: bytes) -> str:
    import logging

    from pypdf import PdfReader

    logging.getLogger("pypdf").setLevel(logging.ERROR)  # malformed files log noisy warnings
    try:
        reader = PdfReader(io.BytesIO(data))
        encrypted = reader.is_encrypted
        pages = [] if encrypted else [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:  # any parser failure on an untrusted file is a bad file
        raise DocumentError("The PDF could not be read.") from exc
    if encrypted:
        raise DocumentError("The PDF is password protected.")
    return "\n".join(pages)


def _tidy(text: str) -> str:
    """Collapse runs of spaces and blank lines; keep one line per paragraph."""
    lines = (re.sub(r"[ \t\u00a0]+", " ", line).strip() for line in text.splitlines())
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()
